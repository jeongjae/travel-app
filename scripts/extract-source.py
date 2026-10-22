"""Read-only selected-source extractor. Never import or execute source modules."""
import argparse
import ast
import hashlib
import json
from datetime import date, datetime, time
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook

MAX_FILE = 64 * 1024 * 1024
MAX_NODES = 50000
MAX_DEPTH = 80


def no_duplicates(pairs):
    out = {}
    for key, value in pairs:
        if key in out:
            raise ValueError(f"duplicate JSON key: {key}")
        out[key] = value
    return out


def json_value(value):
    if isinstance(value, (datetime, date, time)):
        return {"sourceType": type(value).__name__, "isoValue": value.isoformat()}
    return value


def extract(root, spec, repository, commit):
    root = Path(root).resolve()
    result = {"format": "travel-app-source-extract", "version": 1, "sourceRepository": repository,
              "sourceCommit": commit, "adapter": "selected-source-v1", "extractorVersion": "1.0.0",
              "manifest": [], "entries": [], "unresolved": []}
    seen = set()

    def emit(kind, key, locator, raw, issue=None):
        result["entries"].append({"sourceKind": kind, "sourceKey": key, "locator": locator,
                                  "rawValue": raw, "status": "unresolved" if issue else "extracted",
                                  **({"issue": issue} if issue else {})})
        if issue:
            result["unresolved"].append({"sourceKey": key, "reason": issue})

    for selection in spec:
        relative = selection["path"]
        if relative in seen or "\\" in relative or Path(relative).is_absolute():
            raise ValueError("duplicate/unsafe source path")
        seen.add(relative)
        path = (root / relative).resolve()
        if not path.is_relative_to(root) or not path.is_file() or path.stat().st_size > MAX_FILE:
            raise ValueError("source outside root, missing or too large")
        data = path.read_bytes()
        digest = hashlib.sha256(data).hexdigest()
        result["manifest"].append({"path": relative, "byteLength": len(data), "sha256": digest})
        kind = selection["kind"]
        if kind == "python":
            if len(data) > 2 * 1024 * 1024:
                raise ValueError("Python source size limit")
            tree = ast.parse(data.decode("utf-8"), filename=relative)
            stack = [(tree, 0)]
            nodes = 0
            while stack:
                node, depth = stack.pop()
                nodes += 1
                if nodes > MAX_NODES or depth > MAX_DEPTH:
                    raise ValueError("AST resource limit")
                stack.extend((child, depth + 1) for child in ast.iter_child_nodes(node))
            assignments = {}
            for node in tree.body:
                if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
                    assignments.setdefault(node.targets[0].id, []).append(node)
                elif isinstance(node, ast.AnnAssign) and isinstance(node.target, ast.Name):
                    assignments.setdefault(node.target.id, []).append(node)
            for name in selection.get("constants", []):
                candidates = assignments.get(name, [])
                key = f"{relative}:{name}"
                if len(candidates) != 1:
                    emit(kind, key, key, None, "constant missing or assigned more than once")
                    continue
                node = candidates[0]
                locator = f"{relative}:{node.lineno}-{node.end_lineno}:{name}"
                try:
                    # literal_eval rejects calls, names, computed templates and arithmetic.
                    raw = ast.literal_eval(node.value)
                    json.dumps(raw, allow_nan=False)
                    emit(kind, key, locator, raw)
                except (ValueError, TypeError, SyntaxError):
                    emit(kind, key, locator, ast.get_source_segment(data.decode("utf-8"), node.value), "unsupported non-JSON literal/expression")
        elif kind == "xlsx":
            with ZipFile(path) as archive:
                if len(archive.infolist()) > 10000 or sum(e.file_size for e in archive.infolist()) > 128 * 1024 * 1024:
                    raise ValueError("XLSX expansion limit")
            source = load_workbook(path, read_only=True, data_only=False)
            cached = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in source:
                    if sheet.max_row * sheet.max_column > 100000:
                        raise ValueError("XLSX cell limit")
                    cached_rows = cached[sheet.title].iter_rows()
                    for cells, values in zip(sheet.iter_rows(), cached_rows):
                        for cell, cached_cell in zip(cells, values):
                            if cell.value is None:
                                continue
                            key = f"{relative}:{sheet.title}!{cell.coordinate}"
                            raw = {"value": json_value(cell.value), "cellType": cell.data_type,
                                   "cachedValue": json_value(cached_cell.value)}
                            emit(kind, key, key, raw, "formula cache missing" if cell.data_type == "f" and cached_cell.value is None else None)
            finally:
                source.close()
                cached.close()
        elif kind == "json":
            emit(kind, relative, relative, json.loads(data.decode("utf-8"), object_pairs_hook=no_duplicates, parse_constant=lambda x: (_ for _ in ()).throw(ValueError(x))))
        elif kind == "markdown":
            emit(kind, relative, relative, data.decode("utf-8"))
        elif kind == "asset":
            emit(kind, relative, relative, {"path": relative, "sha256": digest, "byteLength": len(data)})
        else:
            emit(kind, relative, relative, None, "unsupported source kind")
        if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
            raise ValueError("source changed while reading")
    return result


def main():
    parser = argparse.ArgumentParser()
    for flag in ["root", "spec", "output", "repository", "commit"]:
        parser.add_argument(f"--{flag}", required=True)
    args = parser.parse_args()
    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"), object_pairs_hook=no_duplicates)
    result = extract(args.root, spec, args.repository, args.commit)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, allow_nan=False, indent=2), encoding="utf-8")
    print(f"files={len(result['manifest'])} entries={len(result['entries'])} unresolved={len(result['unresolved'])}")


if __name__ == "__main__":
    main()
