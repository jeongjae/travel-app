"""P-01 environment/fixture checks only; this is not the P-03 source extractor."""
import ast
import hashlib
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / ".cache" / "pilot-fixtures"


def make_workbook(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Reservations"
    sheet.append(["id", "slug", "date", "schedule", "status", "note", "amount"])
    sheet.append(["booking-xlsx", "shared-source-slug", datetime(2026, 10, 31), "오후", "미확인", "가상 XLSX 예약", 0])
    sheet.append(["booking-formula", "shared-source-slug", None, None, "검토 필요", "수식 캐시 없음", "=1+2"])
    sheet["C2"].number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["F"].width = 30
    book.save(path)
    book.close()


class PilotEnvironment(unittest.TestCase):
    def test_project_python(self):
        self.assertEqual(sys.version_info[:2], (3, 12))
        self.assertEqual(Path(sys.prefix).resolve(), (ROOT / ".venv").resolve())

    def test_ast_source_is_not_executed(self):
        path = ROOT / "tests" / "fixtures" / "booking_literals.py"
        before = hashlib.sha256(path.read_bytes()).digest()
        tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
        assignments = {node.targets[0].id: node.value for node in tree.body if isinstance(node, ast.Assign)}
        records = ast.literal_eval(assignments["PILOT_BOOKINGS"])
        self.assertEqual(records[0]["schedule"], "오후")
        self.assertIsNone(records[0]["action_date"])
        with self.assertRaises(ValueError):
            ast.literal_eval(assignments["UNSUPPORTED_BOOKINGS"])
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)

    def test_xlsx_read_only_formula_and_cache(self):
        CACHE.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=CACHE) as directory:
            path = Path(directory) / "synthetic-reservations.xlsx"
            make_workbook(path)
            before = hashlib.sha256(path.read_bytes()).digest()
            source = load_workbook(path, read_only=True, data_only=False)
            cached = load_workbook(path, read_only=True, data_only=True)
            try:
                self.assertEqual(source.sheetnames, ["Reservations"])
                sheet = source["Reservations"]
                self.assertEqual(sheet["C2"].value, datetime(2026, 10, 31))
                self.assertEqual(sheet["D2"].value, "오후")
                self.assertEqual(sheet["G2"].value, 0)
                self.assertEqual(sheet["G3"].data_type, "f")
                self.assertEqual(sheet["G3"].value, "=1+2")
                self.assertIsNone(cached["Reservations"]["G3"].value)
            finally:
                source.close()
                cached.close()
            self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)


if __name__ == "__main__":
    unittest.main(verbosity=2)
